"""Turn a raw sales CSV into a finished Excel report with one command.

    python auto_report.py sales_raw.csv Sales_Report.xlsx

Steps: load, check the columns, clean the data and log every fix, summarise by
month, branch and category, then write a formatted workbook with a native
Excel chart. The cleaning log is part of the report so the reader can see
exactly what was changed.
"""

import argparse
import sys
import time
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

REQUIRED = ["date", "branch", "category", "units_sold", "unit_price_ngn", "discount", "revenue_ngn"]
NAIRA = '"₦"#,##0'
HEAD_FILL = PatternFill("solid", fgColor="14213D")
HEAD_FONT = Font(bold=True, color="FFFFFF")


def load(path):
    df = pd.read_csv(path)
    missing = [c for c in REQUIRED if c not in df.columns]
    if missing:
        sys.exit(f"Input file is missing columns: {', '.join(missing)}")
    df["date"] = pd.to_datetime(df["date"], errors="coerce")
    return df


def clean(raw):
    """Return the clean frame and a list of (issue, rows affected, what was done)."""
    log = [("Rows in the raw file", len(raw), "")]
    df = raw.copy()

    n = int(df["date"].isna().sum())
    df = df.dropna(subset=["date"])
    log.append(("Rows with an unreadable date", n, "Removed" if n else "None found"))

    n = int(df.duplicated().sum())
    df = df.drop_duplicates()
    log.append(("Exact duplicate rows", n, "Removed" if n else "None found"))

    for col in ("branch", "category"):
        tidy = df[col].astype(str).str.strip().str.title()
        n = int((tidy != df[col]).sum())
        df[col] = tidy
        log.append((f"Inconsistent {col} names (spacing or capitals)", n, "Standardised" if n else "None found"))

    n = int(df["units_sold"].isna().sum())
    rebuild = df["units_sold"].isna() & df["unit_price_ngn"].notna()
    df.loc[rebuild, "units_sold"] = (
        df.loc[rebuild, "revenue_ngn"].abs()
        / (df.loc[rebuild, "unit_price_ngn"] * (1 - df.loc[rebuild, "discount"]))
    ).round()
    still = int(df["units_sold"].isna().sum())
    df = df.dropna(subset=["units_sold"])
    log.append(("Missing unit counts", n,
                f"Rebuilt from revenue and price ({n - still}), removed ({still})" if n else "None found"))

    n = int((df["revenue_ngn"] < 0).sum())
    df["revenue_ngn"] = df["revenue_ngn"].abs()
    log.append(("Negative revenue values", n, "Sign corrected" if n else "None found"))

    df["units_sold"] = df["units_sold"].astype(int)
    log.append(("Rows in the clean data", len(df), ""))
    return df.reset_index(drop=True), log


def summarise(df):
    month = df.groupby(df["date"].dt.to_period("M"))["revenue_ngn"].sum()
    month.index = [p.strftime("%b %Y") for p in month.index]
    branch = df.groupby("branch")["revenue_ngn"].sum().sort_values(ascending=False)
    cat = df.groupby("category")["revenue_ngn"].sum().sort_values(ascending=False)
    return month, branch, cat


def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        cell.alignment = Alignment(horizontal="left" if c == 1 else "right", vertical="center")


def widths(ws, spec):
    for i, w in enumerate(spec, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def write_table(ws, title, series, value_header):
    ws["A1"], ws["A1"].font = title, Font(bold=True, size=14)
    ws.append([])
    ws.append([title.split(" by ")[-1].capitalize(), value_header, "Share"])
    style_header(ws, 3, 3)
    total = float(series.sum())
    for name, value in series.items():
        ws.append([name, float(value), float(value) / total])
    last = ws.max_row
    for r in range(4, last + 1):
        ws.cell(r, 2).number_format = NAIRA
        ws.cell(r, 3).number_format = "0.0%"
    ws.freeze_panes = "A4"
    widths(ws, [24, 20, 10])
    return last


def add_chart(ws, last, title, anchor):
    ch = BarChart()
    ch.type, ch.title, ch.style = "col", title, 10
    ch.legend = None
    ch.y_axis.title, ch.y_axis.numFmt = "Revenue (₦)", '#,##0'
    ch.add_data(Reference(ws, min_col=2, min_row=3, max_row=last), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=1, min_row=4, max_row=last))
    ch.height, ch.width = 9, 20
    ws.add_chart(ch, anchor)


def build_workbook(df, log, out):
    month, branch, cat = summarise(df)
    wb = Workbook()

    s = wb.active
    s.title = "Summary"
    s["A1"], s["A1"].font = "Sales report", Font(bold=True, size=16)
    s["A2"] = f"Period: {df['date'].min():%d %b %Y} to {df['date'].max():%d %b %Y}"
    rows = [
        ("Total revenue", float(df["revenue_ngn"].sum()), NAIRA),
        ("Units sold", int(df["units_sold"].sum()), "#,##0"),
        ("Average revenue per day", float(df.groupby("date")["revenue_ngn"].sum().mean()), NAIRA),
        ("Best month", f"{month.idxmax()} ({month.max() / 1e6:,.0f}M)", None),
        ("Weakest month", f"{month.idxmin()} ({month.min() / 1e6:,.0f}M)", None),
        ("Top branch", f"{branch.index[0]} ({branch.iloc[0] / branch.sum():.1%})", None),
        ("Lowest branch", f"{branch.index[-1]} ({branch.iloc[-1] / branch.sum():.1%})", None),
        ("Top category", f"{cat.index[0]} ({cat.iloc[0] / cat.sum():.1%})", None),
    ]
    s.append([])
    s.append(["Measure", "Value"])
    style_header(s, 4, 2)
    for label, value, fmt in rows:
        s.append([label, value])
        if fmt:
            s.cell(s.max_row, 2).number_format = fmt
        s.cell(s.max_row, 2).alignment = Alignment(horizontal="right")
    widths(s, [30, 34])

    m = wb.create_sheet("By month")
    last = write_table(m, "Revenue by month", month, "Revenue")
    add_chart(m, last, "Revenue by month", "E3")
    write_table(wb.create_sheet("By branch"), "Revenue by branch", branch, "Revenue")
    write_table(wb.create_sheet("By category"), "Revenue by category", cat, "Revenue")

    lg = wb.create_sheet("Cleaning log")
    lg.append(["Check", "Rows", "Action"])
    style_header(lg, 1, 3)
    for issue, n, action in log:
        lg.append([issue, n, action])
    for r in range(2, lg.max_row + 1):
        lg.cell(r, 2).number_format = "#,##0"
    widths(lg, [50, 12, 56])

    d = wb.create_sheet("Clean data")
    cols = list(df.columns)
    d.append(cols)
    style_header(d, 1, len(cols))
    for rec in df.itertuples(index=False):
        d.append([v.to_pydatetime() if isinstance(v, pd.Timestamp) else v for v in rec])
    for r in range(2, d.max_row + 1):
        d.cell(r, 1).number_format = "yyyy-mm-dd"
        d.cell(r, cols.index("revenue_ngn") + 1).number_format = '#,##0.00'
    d.freeze_panes = "A2"
    d.auto_filter.ref = d.dimensions
    widths(d, [12, 18, 16, 12, 16, 10, 16])

    wb.save(out)
    return len(wb.sheetnames)


def main():
    ap = argparse.ArgumentParser(description="Clean a sales CSV and write an Excel report.")
    ap.add_argument("input", help="path to the raw sales CSV")
    ap.add_argument("output", nargs="?", default="Sales_Report.xlsx", help="Excel file to write")
    args = ap.parse_args()

    t0 = time.time()
    raw = load(args.input)
    df, log = clean(raw)
    n_sheets = build_workbook(df, log, args.output)

    print(f"Read      {len(raw):>7,} rows from {Path(args.input).name}")
    for issue, n, action in log[1:-1]:
        if n:
            print(f"Fixed     {n:>7,}  {issue.lower()}: {action.lower()}")
    print(f"Clean     {len(df):>7,} rows")
    print(f"Wrote     {args.output}  ({n_sheets} sheets, 1 chart) in {time.time() - t0:.1f}s")


if __name__ == "__main__":
    main()
